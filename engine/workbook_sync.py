"""Populate all 5 sheets of the working submission workbook copy with verified evidence.

Never modifies inputs/official_submission_workbook.xlsx. Preserves all template
formulas, merged ranges, freeze panes, data validations, and candidate-entered inputs.

Run: py -3.12 -m engine.workbook_sync
"""
from __future__ import annotations

import copy
import datetime
import hashlib
from pathlib import Path
from typing import Any

import openpyxl

from .workbook_population import populate_workbook_copy

EXPECTED_OFFICIAL_SHA256 = "9142119c5633480b40e9de997883265cc930e4b868f2d128845fa9677b321b9a"

DEFAULT_START_VALUES: dict[str, Any] = {
    "B5": None,
    "B6": "Viswanath V M",
    "B7": "23i275@psgtech.ac.in",
    "B8": "PSG College of Technology (B.Tech-IT)",
    "B9": datetime.datetime(2026, 10, 19, 0, 0),
    "B10": "PENDING FINAL CANDIDATE SUBMISSION",
    "B11": "submission-package/main-output/ (calendar.csv + content/exports/ + content/designs/)",
    "B12": "submission-package/engine/ (root repository with README.md, engine/, config/, tests/)",
    "B13": "launch-kit/launch-plan.md (and submission-package/launch-kit/)",
    "B14": "PENDING CANDIDATE RECORDING - script at launch-kit/demo-1-product-walkthrough.md",
    "B15": "PENDING CANDIDATE RECORDING - script at launch-kit/demo-2-technical-workflow.md",
    "B21": "No",
    "F5": "No",
    "F6": "Yes",
    "F7": "Yes",
    "F8": "Yes",
    "F9": "No",
    "F10": "No",
}

PLACEHOLDER_PREFIXES = (
    "UNKNOWN",
    "Viswanath (candidate to confirm",
    "2026-10-19 (earliest audit record",
    "PENDING FINAL CANDIDATE SUBMISSION",
    "PENDING CANDIDATE RECORDING",
)


def _is_placeholder(val: Any) -> bool:
    if val is None:
        return True
    if isinstance(val, str):
        s = val.strip()
        if not s or s.startswith(PLACEHOLDER_PREFIXES):
            return True
    return False


def _harvest_candidate_inputs(candidates: list[Path], official: Path) -> dict[tuple[str, str], Any]:
    """Read legitimate candidate-entered non-formula values from existing working copies."""
    wo = openpyxl.load_workbook(official, data_only=False)
    formula_coords: set[tuple[str, str]] = set()
    for ws in wo.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_coords.add((ws.title, cell.coordinate))
    wo.close()

    harvested: dict[tuple[str, str], Any] = {}
    watch_coords = [
        ("00_START", c)
        for c in ("B5", "B6", "B7", "B8", "B9", "B10", "B11", "B12", "B13", "B14", "B15", "B21",
                  "F5", "F6", "F7", "F8", "F9", "F10")
    ]
    for path in candidates:
        if not path.is_file():
            continue
        try:
            wb = openpyxl.load_workbook(path, data_only=False)
        except Exception:
            continue
        try:
            for sheet_name, coord in watch_coords:
                if (sheet_name, coord) in formula_coords:
                    continue
                val = wb[sheet_name][coord].value
                if isinstance(val, str) and val.startswith("="):
                    continue
                if coord in ("B5", "B6", "B7", "B8", "B9", "B10", "B14", "B15"):
                    if not _is_placeholder(val):
                        harvested[(sheet_name, coord)] = val
                elif coord in ("B21", "F5", "F6", "F7", "F8", "F9", "F10"):
                    if val == "Yes":
                        harvested[(sheet_name, coord)] = val
        finally:
            wb.close()
    return harvested


def _restore_official_formulas_and_validations(wb: openpyxl.Workbook, official_path: Path) -> None:
    """Guarantee that every official formula and data-validation rule is preserved verbatim."""
    wo = openpyxl.load_workbook(official_path, data_only=False)
    try:
        for ws_off in wo.worksheets:
            ws_dst = wb[ws_off.title]
            for row in ws_off.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        ws_dst[cell.coordinate].value = cell.value
            ws_dst.data_validations.dataValidation = [
                copy.deepcopy(dv) for dv in ws_off.data_validations.dataValidation
            ]
    finally:
        wo.close()


def sync_working_workbooks(root: str | Path = ".") -> list[str]:
    root = Path(root)
    official = root / "inputs" / "official_submission_workbook.xlsx"
    calendar = (root / "calendar_UPDATED.csv") if (root / "calendar_UPDATED.csv").exists() else (root / "calendar.csv")
    schema = root / "config" / "workbook-schema.json"

    before_hash = hashlib.sha256(official.read_bytes()).hexdigest().lower()
    if before_hash != EXPECTED_OFFICIAL_SHA256:
        raise ValueError(f"Official workbook SHA-256 mismatch before sync: {before_hash}")

    draft_path = root / "outputs" / "IncuBrix_Final_Submission_DRAFT.xlsx"
    working_path = root / "submission-package" / "workbook" / "IncuBrix_Option2_Submission_Workbook_WORKING_COPY.xlsx"
    updated_path = root / "submission-package" / "workbook" / "IncuBrix_Option2_Submission_Workbook_WORKING_COPY_UPDATED.xlsx"

    harvested = _harvest_candidate_inputs([updated_path, working_path, draft_path], official)

    targets = [draft_path, working_path]
    if updated_path.exists():
        targets.append(updated_path)

    updated: list[str] = []
    for dest in targets:
        dest.parent.mkdir(parents=True, exist_ok=True)
        actual_dest = dest
        try:
            dest.unlink(missing_ok=True)
        except PermissionError:
            if dest == working_path:
                actual_dest = updated_path
                try:
                    actual_dest.unlink(missing_ok=True)
                except PermissionError:
                    updated.append(f"{actual_dest} (locked by spreadsheet viewer)")
                    continue
            else:
                updated.append(f"{dest} (locked by spreadsheet viewer)")
                continue

        populate_workbook_copy(
            workbook_path=official,
            csv_path=calendar,
            schema_path=schema,
            output_path=actual_dest,
        )
        wb = openpyxl.load_workbook(actual_dest, data_only=False)
        try:
            _populate_non_asset_tabs(wb, harvested)
            _restore_official_formulas_and_validations(wb, official)
            wb.save(actual_dest)
        finally:
            wb.close()
        updated.append(str(actual_dest))

    after_hash = hashlib.sha256(official.read_bytes()).hexdigest().lower()
    if after_hash != EXPECTED_OFFICIAL_SHA256:
        raise ValueError(f"Official workbook SHA-256 changed after sync: {after_hash}")

    return updated


def _populate_non_asset_tabs(wb: openpyxl.Workbook, harvested: dict[tuple[str, str], Any] | None = None) -> None:
    harvested = harvested or {}
    # 00_START
    s0 = wb["00_START"]
    for coord, default_val in DEFAULT_START_VALUES.items():
        s0[coord] = harvested.get(("00_START", coord), default_val)

    s0["B16"] = (
        "IncuBrix Web App (app.incubrix.com free tier, 300 credits, AI Script Assistant); "
        "OpenRouter Free Tier (openrouter/free: cohere/north-mini-code:free & nvidia/nemotron-3-ultra-550b-a55b:free "
        "in saved acceptance runs); Python 3.12 local content engine (mock deterministic provider); "
        "HTML/CSS + Headless Chrome/Edge (export-instagram.ps1); FFmpeg/ffprobe (video inspection & 720x1280 scaling); "
        "Synthetic presenter video generator (Gemini/Veo 4-point star watermark in clean uncaptioned YT-01..10); "
        "Public LinkedIn pages (Hootsuite, Buffer, Later, Metricool, Sprout Social)."
    )
    s0["B17"] = "$0.00 (all tools, local Python engine, native schedulers, and OpenRouter free-tier models used at $0 cost)"
    s0["B18"] = (
        "Candidate logged into app.incubrix.com and captured 7 product screenshots; authored/reviewed 30 asset drafts, "
        "HTML slide templates, and video scripts; built Python verification/audit/pipeline modules (350 tests); "
        "human rights sign-off, UTM destination confirmation, 20-person audience test collection, and 2 unedited "
        "video recordings remain for candidate completion."
    )

    # 01_INSIGHT
    s1 = wb["01_INSIGHT"]
    obs = [
        (
            "Home Dashboard & Credit-Metered Entry (Text, Audio, Video, Clips, Editor)",
            "Logged into app.incubrix.com (Viswanath, 300 credits free tier). Dashboard presents a single topic/script prompt bar, 5 mode icons (Text, Audio, Video, Clips, Editor), and Popular templates (A-SERIES..E-SERIES).",
            "Worked well: Clean, low-cognitive-load entry with immediate 300-credit transparency. Friction: Template cards (A-SERIES..E-SERIES) do not preview target platform, format, or credit cost before clicking.",
            "inputs/product-screenshots/01-dashboard-home-20261025-010726.png",
            "Showing creators how one workspace turns a raw topic into structured assets is a clear top-of-funnel hook, and clarifying template previews can improve first-session activation.",
        ),
        (
            "AI Script Assistant Cold-Start & Trending Topic Prompts",
            "Opening the Text tool launched AI Script Assistant with 5 pre-populated Ideas from Trending Topics plus a Prompt Idea helper and freeform input.",
            "Worked well: Eliminates blank-page paralysis with immediate topic starters. Friction: Default trending topics skew toward macro business/finance news unless a custom niche topic is entered.",
            "inputs/product-screenshots/02-ai-script-assistant-home-20261025-010835.png",
            "Solving daily blank-page fatigue is a primary creator pain point, and social content can contrast generic blank-box AI prompts with guided topic ideation.",
        ),
        (
            "Four-Step Guided Intake Before Drafting (Format -> Structure -> Treatment -> Personal Angle)",
            "Submitted prompt 'Five practical steps to plan a week of social media content'. Assistant walked through 4 choices: FORMAT (LinkedIn post), STRUCTURE (Bullet points), TREATMENT (Educational explainer), and optional PERSONAL ANGLE.",
            "Worked well: Fast chip-based constraints lock in format, structure, tone, and personal proof before spending credits on full drafts. Friction: Requires 4 sequential clicks before seeing angles.",
            "inputs/product-screenshots/03-guided-step1-format-20261025-010941.png",
            "Directly supports anti-slop creator positioning: showing how 4 quick constraints turn 1 rough idea into structured copy makes high-converting educational demos.",
        ),
        (
            "Multi-Angle Direction Cards with Readiness Triage (Ready to draft vs Needs one detail)",
            "Clicking Generate directions produced 3 distinct LINKEDIN POST angle cards (PLAYBOOK, OBSERVATION, PREDICTION) plus Show 3 more angles, with status badges Ready to draft (green) vs Needs one detail (amber).",
            "Worked well: Lets creators compare distinct hooks before drafting and flags when an angle needs a real detail instead of hallucinating facts. Friction: Needs one detail badge does not inline-specify the missing detail until expanded.",
            "inputs/product-screenshots/07-generated-angles-output-20261025-011742.png",
            "'1 Idea -> 3 Angles (Playbook, Observation, Prediction)' is a repeatable framework taught across our LinkedIn/IG carousels and YouTube Shorts.",
        ),
        (
            "Gated Post-Draft Utilities (Draft Actions) & Coming Soon Collaboration Panel (Visible UI)",
            "Right sidebar displays PRIMARY utilities (Publishing Readiness, BrixScore, Humaniser) and ADVANCED utilities (Series Builder, Platform Pack, Engagement Kit, Angle Lab) gated behind 'Generate a draft to unlock these actions', plus Collaborators/Comments marked 'Coming soon'.",
            "Worked well: Surfaces quality/repurposing utilities and honestly labels team collaboration as Coming soon. Scope note: Observed as visible UI elements in screenshot 02 (post-draft execution was not captured in this screenshot set).",
            "inputs/product-screenshots/02-ai-script-assistant-home-20261025-010835.png",
            "Highlights quality-first roadmap (BrixScore, Humaniser, Platform Pack) while ensuring marketing never advertises Coming soon features as live.",
        ),
    ]
    for idx, (area, what, diff, evid, why) in enumerate(obs, start=6):
        s1.cell(idx, 2).value = area
        s1.cell(idx, 3).value = what
        s1.cell(idx, 4).value = diff
        s1.cell(idx, 5).value = evid
        s1.cell(idx, 6).value = why

    opps = [
        (
            "Own the '1 Idea -> 3 Structured Angles' Anti-Slop Workflow Category",
            "Product screenshots 03..07 (AI Script Assistant 4-step guided intake -> Playbook, Observation, Prediction cards) + inputs/competitor-audit.md (incumbents focus on downstream scheduling queues).",
            "High save rates and qualified free-tier sign-ups (300 credits) from creators frustrated by generic one-shot AI prompts.",
            "Carousel save rate, Short 3-second view rate, UTM link CTR (utm_campaign=incubrix_content_engine), and first AI Script Assistant angle generated.",
            "PROPOSED 30-DAY TARGET (baseline=0): 30 posts published; 8,000-15,000 views/impressions; >=3.5% LI/IG engagement rate; 120-250 UTM clicks; 15-30 free sign-ups.",
        ),
        (
            "Teach Repeatable '1 Long-Form Source -> 1 Full Week Across LinkedIn, Instagram & Shorts' Systems",
            "Product screenshots 01..02 (Text, Audio, Video, Clips, Editor + Series Builder / Platform Pack UI) + engine/pipeline.py (1 brief -> 12 platform-checked assets in <1s mock / <30m).",
            "Positions IncuBrix as the weekly operating system for podcasters, YouTubers, and small editing teams (1-3 people).",
            "Average percentage viewed on YT-01..10, carousel saves on LI-01..07 & IG-01..08, and UTM clicks on repurposing assets (LI-02, LI-07, IG-02, IG-06, YT-02, YT-06).",
            "PROPOSED 30-DAY TARGET (baseline=0): >=10 repurposing assets published; >=40% avg view duration on 16-24s Shorts; 50+ saves on repurposing carousels.",
        ),
        (
            "Build Creator Trust via Transparent Pre-Publish Quality, Rights & Readiness Checklists",
            "Product screenshots 02 & 07 (Publishing Readiness, BrixScore, Humaniser, Needs one detail badge, honest Coming soon labels) + outputs/asset-audit/asset-audit.md.",
            "Differentiates IncuBrix from one-click AI spam tools among brand-conscious creators and editors who require accuracy and rights safety.",
            "Save-to-impression ratio on checklist assets (LI-06, IG-05, IG-09, YT-05) and Option B 20-person audience preference test win rate.",
            "PROPOSED 30-DAY TARGET (baseline=0): 100% published assets pass pre-publish audit (0 rights/claim issues); 30-50 checklist saves; >=60% preference in 20-person test.",
        ),
    ]
    for idx, (opp, evid, imp, meas, targ) in enumerate(opps, start=14):
        s1.cell(idx, 2).value = opp
        s1.cell(idx, 3).value = evid
        s1.cell(idx, 4).value = imp
        s1.cell(idx, 5).value = meas
        s1.cell(idx, 6).value = targ

    # 03_ENGINE_LAUNCH
    s3 = wb["03_ENGINE_LAUNCH"]
    proofs = [
        ("Yes", "README.md", "py -3.12 -m engine pipeline --input .\\inputs\\sample.md --source-id my-source-001"),
        ("Yes", "pyproject.toml", "Runs offline on a fresh Python 3.12 install using deterministic mock provider; 350 pytest tests pass."),
        ("Yes", "engine/pipeline.py", "Blocks exact and near-duplicates (>=0.85 similarity) within batch and across runs via data/pipeline-ledger.json."),
        ("Yes", "engine/asset_audit.py", "Checks platform limits, required fields, prohibited phrases, and factual/performance claims before export."),
        ("Yes", "engine/workbook_population.py", "Exports exact 23-column official schema to CSV and populates 02_ASSETS preserving V/W formulas."),
        ("Yes", "outputs/final-acceptance-20261025/run-20261025T035339Z/run-record.json", "--refresh updates existing items for source_id in place (items_updated=12, cumulative_rows=12) without duplicate IDs."),
        ("Yes", "outputs/final-acceptance-20261025/run-20261025T035344Z/run-record.json", "Structured JSONL log records batch_attempt_failed (PROVIDER_ERROR) and recovers automatically on retry (batch_recovered)."),
        ("Yes", "tests/ (py -3.12 -m pytest -q)", "350 automated tests passing in Python 3.12 (including pipeline, refresh, duplicate, recovery, audit, package_verify, and workbook tests)."),
        ("Yes", ".gitignore", ".env is gitignored and never copied to submission-package/; repository contains zero secret keys."),
    ]
    for idx, (done, evid, note) in enumerate(proofs, start=6):
        s3.cell(idx, 2).value = done
        s3.cell(idx, 3).value = evid
        s3.cell(idx, 8).value = note

    runs = [
        ("184 chars (inputs/sample.md)", 12, 0, 0, "0.0009 min (0.054s mock)", "$0.00 (mock; saved openrouter/free runs also $0.00)", "outputs/final-acceptance-20261025/run-20261025T035335Z/run-record.json"),
        ("184 chars (inputs/sample.md)", 12, 0, 12, "0.0009 min (0.055s mock)", "$0.00 (mock)", "outputs/final-acceptance-20261025/run-20261025T035337Z/run-record.json (items_exported=0, duplicates_removed=12)"),
        ("1202 chars (inputs/sample-b.md --refresh)", 12, 0, 0, "0.0011 min (0.064s mock)", "$0.00 (mock)", "outputs/final-acceptance-20261025/run-20261025T035339Z/run-record.json (items_updated=12, cumulative_rows=12)"),
        ("184 chars (inputs/sample.md --simulate-failure-once)", 12, 0, 0, "0.0021 min (0.128s mock)", "$0.00 (mock)", "outputs/final-acceptance-20261025/run-20261025T035344Z/run-record.json (batch_attempt_failed -> batch_recovered)"),
    ]
    for idx, (inp, prod, rej, dup, rt, cost, ev) in enumerate(runs, start=18):
        s3.cell(idx, 2).value = inp
        s3.cell(idx, 3).value = prod
        s3.cell(idx, 4).value = rej
        s3.cell(idx, 5).value = dup
        s3.cell(idx, 6).value = rt
        s3.cell(idx, 7).value = cost
        s3.cell(idx, 8).value = ev

    s3["B24"] = "Primary: Podcasters, YouTubers, and expert-led creators. Secondary: Freelance editors and small content teams (1-3 people) in US, CA, UK, IE, AU, NZ, and SG."
    s3["B25"] = "LinkedIn (10 assets: 9 carousels + 1 graphic for B2B/expert creators & outbound CTR), Instagram (10 assets: 8 carousels + 2 graphics for visual workflow saves), YouTube Shorts (10 clean uncaptioned vertical videos 16s-24s for video-first creator discovery)."
    s3["B26"] = "Owner: Growth Engineer + Reviewer (15m Monday sign-off). Tools: Python engine + native schedulers. Volume: 7 assets (LI-01..03, IG-01..02, YT-01..02) after confirming UTM base URL and rights sign-off."
    s3["B27"] = "Publish 7-8 assets/week (Days 8-30: LI-04..10, IG-03..10, YT-03..10); every Monday run engine pipeline --refresh on 1 new source to generate 12 next-week drafts in <30 min."
    s3["B28"] = "$0.00/month required stack: Python 3.12 engine + mock/openrouter-free ($0), HTML/CSS + headless browser PNG export ($0), FFmpeg/ffprobe video verification ($0), native LinkedIn/Meta/YouTube schedulers ($0). Optional $150-$300 paid boost in Weeks 3-4."
    s3["B29"] = "PROPOSED 30-DAY TARGETS (current baseline = 0): 30 posts published; 8,000-15,000 impressions/views; >=3.5% LI/IG engagement rate (>=40% YT avg view duration); 80-150 saves; 120-250 tracked UTM clicks; 15-30 free-tier sign-ups."
    s3["B30"] = "Risk: Unverified AI product/performance claims or un-cleared media/synthetic video disclosures slipping into live posts. Mitigation: Automated claim/brand audit (engine/asset_audit.py) + mandatory human sign-off on launch-kit/rights-inventory.csv + YouTube synthetic content label."
    s3["B31"] = "1) Confirm approved landing page URL for the 30 UTM links in outputs/proposed-utm-links-UNVERIFIED.csv; 2) Scheduler access to official LinkedIn, Instagram, and YouTube accounts; 3) 15-min human sign-off on launch-kit/rights-inventory.csv (logo & YT synthetic presenter); 4) Optional boost budget decision ($0 default)."

    # 04_SELF_CHECK
    s4 = wb["04_SELF_CHECK"]
    s4["C6"] = 0
    s4["E6"] = "outputs/asset-audit/asset-audit.md (30 assets created and exported, 30/30 pass automated checks, 0 marked Ready until human rights and UTM sign-off)"
    s4["C7"] = 3
    s4["E7"] = "calendar.csv (LinkedIn: 10, Instagram: 10, YouTube Shorts: 10)"
    s4["C8"] = 0
    s4["E8"] = "content/exports/youtube-shorts/ (10 clean uncaptioned vertical 720x1280 videos 16.0s-24.0s <=60s exist on disk, 0 marked Ready in 02_ASSETS until human rights/AI-label sign-off)"
    s4["C9"] = 0
    s4["E9"] = "content/exports/ (20 carousels/graphics exist on disk as PNGs: 17 carousels + 3 graphics, 0 marked Ready in 02_ASSETS until human sign-off)"
    s4["C10"] = 30
    s4["E10"] = "launch-kit/rights-inventory.csv (D10 counts rows where rights_checked='No', all 30 remain 'No' / PENDING HUMAN REVIEW until candidate verifies licences)"
    s4["C11"] = "Yes"
    s4["E11"] = "outputs/final-acceptance-20261025/run-20261025T035335Z/run-record.json (12 assets across 3 platforms in 0.054s mock, saved live openrouter/free runs in outputs/acceptance-live-10-20261023/)"
    s4["C12"] = "No"
    s4["E12"] = "launch-kit/audience-test-plan.md (Option B 20-person preference test kit and CLI summarizer ready, 0 real participant responses collected yet)"

    for r_idx, note in [
        (17, ("No", "30 draft assets (10 videos)", "outputs/asset-audit/asset-audit.md")),
        (18, ("No", "0 participants tested so far", "launch-kit/audience-test-plan.md")),
        (19, ("No", "21 assets/run verified (--per-platform 7)", "outputs/final-acceptance-20261025/scale/run-20261025T035346Z/run-record.json")),
        (20, ("No", "Not implemented in initial 30-day cycle", "inputs/initial-content-strategy.md")),
    ]:
        s4.cell(r_idx, 2).value = note[0]
        s4.cell(r_idx, 3).value = note[1]
        s4.cell(r_idx, 5).value = note[2]

    scores = [
        (24, 9, "5 product observations backed by 7 timestamped screenshots from app.incubrix.com (inputs/product-screenshots/01..07), 3 opportunities (outputs/social-media-opportunities.md), and 5-competitor LinkedIn audit (inputs/competitor-audit.md); post-draft locked utilities honestly distinguished."),
        (25, 25, "Complete strategy (inputs/initial-content-strategy.md) + 30 assets across LinkedIn, Instagram & YouTube Shorts (10 clean uncaptioned 720x1280 videos 16s-24s, 20 carousels/graphics) with canonical sources; 30/30 pass automated audit (20/20 = 100% random sample). Deducted 10 pts because 20-person audience test and final Ready/UTM sign-off require candidate completion."),
        (26, 20, "One-command pipeline (py -3.12 -m engine pipeline) in README produces 12-21 platform-specific assets from 1 input, enforces brand/claim rules, removes duplicates via ledger, supports --refresh, recovers from failure with JSONL logs, and passes 350 pytest tests."),
        (27, 9, "Actionable 30-day calendar mapping all 30 assets, 12-stage daily/weekly workflow, owners, $0 required tool stack, proposed 30-day targets, and risk mitigations in launch-kit/launch-plan.md; awaits Day-1 UTM base URL approval."),
        (28, 20, "Full evidence traceability across all 30 assets (content/canonical-sources.csv, outputs/asset-audit/asset-audit.md, launch-kit/rights-inventory.csv) and reproducible CLI commands; 2 unedited video recordings remain for candidate to record using launch-kit/demo-1 & demo-2 scripts."),
    ]
    for r_idx, sc, why in scores:
        s4.cell(r_idx, 3).value = sc
        s4.cell(r_idx, 5).value = why


if __name__ == "__main__":
    paths = sync_working_workbooks(Path(__file__).resolve().parent.parent)
    for p in paths:
        print(f"Synced workbook: {p}")
