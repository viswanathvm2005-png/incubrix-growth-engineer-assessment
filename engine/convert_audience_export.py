import csv
from pathlib import Path
from openpyxl import load_workbook

src = Path("launch-kit/audience-test-raw-export-latest.xlsx")
dst = Path("launch-kit/audience-test-raw-export-working.csv")

if not src.exists():
    raise FileNotFoundError(f"Source file not found: {src}")

wb = load_workbook(src, read_only=True, data_only=True)

try:
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)

    if headers is None:
        raise ValueError("The Excel worksheet is empty.")

    count = 0

    with dst.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["" if v is None else str(v) for v in headers])

        for row in rows:
            if row and any(v is not None and str(v).strip() for v in row):
                writer.writerow(["" if v is None else str(v) for v in row])
                count += 1

    print("Source:", src)
    print("Temporary CSV:", dst)
    print("Data rows exported:", count)
    print("CSV size:", dst.stat().st_size, "bytes")
finally:
    wb.close()
