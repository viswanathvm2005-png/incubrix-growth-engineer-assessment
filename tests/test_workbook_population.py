import csv
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from engine.workbook_population import (
    MAX_ASSETS,
    WorkbookPopulationError,
    populate_workbook_copy,
)


ROOT = Path(__file__).resolve().parents[1]
OFFICIAL = ROOT / "inputs" / "official_submission_workbook.xlsx"
SCHEMA = ROOT / "config" / "workbook-schema.json"


class WorkbookPopulationTests(unittest.TestCase):
    def setUp(self):
        if not OFFICIAL.is_file():
            self.skipTest("Local official workbook is not available.")
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.schema = json.loads(SCHEMA.read_text(encoding="utf-8-sig"))
        self.columns = self.schema["columns"]

    def make_csv(self, rows, *, name="calendar.csv", columns=None):
        path = self.directory / name
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=columns or self.columns,
                lineterminator="\n",
            )
            writer.writeheader()
            writer.writerows(rows)
        return path

    def blank_row(self, asset_id="test-asset-001"):
        row = {column: "" for column in self.columns}
        row.update({
            "asset_id": asset_id,
            "platform": "LinkedIn",
            "format": "Text post",
            "status": "Draft",
        })
        return row

    def populate(self, csv_path, output_path):
        return populate_workbook_copy(
            workbook_path=OFFICIAL,
            csv_path=csv_path,
            schema_path=SCHEMA,
            output_path=output_path,
        )

    def test_creates_copy_and_preserves_formulas_validations_and_sheets(self):
        csv_path = self.make_csv([self.blank_row()])
        output = self.directory / "populated.xlsx"

        result = self.populate(csv_path, output)

        self.assertEqual(result["status"], "CREATED")
        self.assertFalse(result["original_workbook_modified"])
        wb = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(wb.close)
        ws = wb["02_ASSETS"]
        self.assertEqual(wb.sheetnames, [
            "00_START", "01_INSIGHT", "02_ASSETS",
            "03_ENGINE_LAUNCH", "04_SELF_CHECK",
        ])
        self.assertEqual(ws["A5"].value, "test-asset-001")
        self.assertEqual(ws["B5"].value, "LinkedIn")
        self.assertEqual(ws["C5"].value, "Text post")
        self.assertTrue(ws["V5"].data_type == "f")
        self.assertTrue(ws["W5"].data_type == "f")
        self.assertEqual(len(ws.data_validations.dataValidation), 5)
        self.assertEqual(ws.freeze_panes, "C5")

    def test_rejects_existing_output_without_overwriting(self):
        csv_path = self.make_csv([self.blank_row()])
        output = self.directory / "existing.xlsx"
        output.write_bytes(b"keep this file")
        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, output)
        self.assertEqual(output.read_bytes(), b"keep this file")

    def test_rejects_source_as_output(self):
        csv_path = self.make_csv([self.blank_row()])
        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, OFFICIAL)

    def test_rejects_header_order_mismatch(self):
        bad_columns = list(self.columns)
        bad_columns[0], bad_columns[1] = bad_columns[1], bad_columns[0]
        csv_path = self.make_csv(
            [self.blank_row()], name="bad.csv", columns=bad_columns
        )
        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, self.directory / "bad.xlsx")

    def test_rejects_invalid_dropdown_value(self):
        row = self.blank_row()
        row["platform"] = "Made-up Platform"
        csv_path = self.make_csv([row])
        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, self.directory / "bad-value.xlsx")

    def test_rejects_duplicate_asset_ids(self):
        first = self.blank_row("duplicate-asset-001")
        second = self.blank_row("duplicate-asset-001")
        csv_path = self.make_csv([first, second])
        output = self.directory / "duplicate-ids.xlsx"

        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, output)

        self.assertFalse(output.exists())

    def test_rejects_blank_asset_id(self):
        row = self.blank_row("")
        csv_path = self.make_csv([row])
        output = self.directory / "blank-id.xlsx"

        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, output)

        self.assertFalse(output.exists())

    def test_rejects_more_than_150_assets(self):
        rows = [
            self.blank_row(f"asset-{index:03}")
            for index in range(MAX_ASSETS + 1)
        ]
        csv_path = self.make_csv(rows)
        with self.assertRaises(WorkbookPopulationError):
            self.populate(csv_path, self.directory / "too-many.xlsx")

    def test_preserves_empty_unknown_values(self):
        row = self.blank_row()
        row["final_file_url"] = ""
        row["source_file_url"] = ""
        csv_path = self.make_csv([row])
        output = self.directory / "unknowns.xlsx"
        self.populate(csv_path, output)
        wb = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(wb.close)
        ws = wb["02_ASSETS"]
        self.assertIsNone(ws["I5"].value)
        self.assertIsNone(ws["J5"].value)

    def test_populated_headers_match_exact_schema_order(self):
        csv_path = self.make_csv([self.blank_row()])
        output = self.directory / "headers.xlsx"
        self.populate(csv_path, output)

        wb = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(wb.close)
        ws = wb["02_ASSETS"]
        actual = [ws.cell(4, col).value for col in range(1, 24)]

        self.assertEqual(len(self.columns), 23)
        self.assertEqual(actual, self.columns)

    def test_all_csv_fields_map_to_their_excel_columns(self):
        row = self.blank_row("mapping-test-001")
        row.update({
            "platform": "Instagram",
            "format": "Carousel",
            "content_pillar": "Educational",
            "hook": "A test hook",
            "caption": "A test caption",
            "cta": "Learn more",
            "source_input": "UNKNOWN - needs input from candidate",
            "tools_used": "Prototype test",
            "rights_checked": "No",
            "ai_label_needed": "Yes",
            "status": "Draft",
            "notes": "Mapping test only",
        })
        csv_path = self.make_csv([row])
        output = self.directory / "mapping.xlsx"
        self.populate(csv_path, output)

        wb = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(wb.close)
        ws = wb["02_ASSETS"]

        for column_index, header in enumerate(self.columns[:21], start=1):
            expected = row[header] or None
            self.assertEqual(
                ws.cell(5, column_index).value,
                expected,
                f"Incorrect mapping for {header}",
            )

        self.assertEqual(ws["V5"].data_type, "f")
        self.assertEqual(ws["W5"].data_type, "f")

    def test_multiple_assets_use_consecutive_rows_and_preserve_formulas(self):
        first = self.blank_row("multi-asset-001")
        second = self.blank_row("multi-asset-002")
        second["platform"] = "Instagram"
        second["format"] = "Carousel"
        csv_path = self.make_csv([first, second])
        output = self.directory / "multiple.xlsx"

        result = self.populate(csv_path, output)

        self.assertEqual(result["assets_written"], 2)
        wb = openpyxl.load_workbook(output, data_only=False)
        self.addCleanup(wb.close)
        ws = wb["02_ASSETS"]

        self.assertEqual(ws["A5"].value, "multi-asset-001")
        self.assertEqual(ws["A6"].value, "multi-asset-002")
        self.assertEqual(ws["B6"].value, "Instagram")
        self.assertEqual(ws["C6"].value, "Carousel")
        self.assertEqual(ws["V5"].data_type, "f")
        self.assertEqual(ws["W5"].data_type, "f")
        self.assertEqual(ws["V6"].data_type, "f")
        self.assertEqual(ws["W6"].data_type, "f")


if __name__ == "__main__":
    unittest.main()
