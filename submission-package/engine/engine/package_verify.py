"""Independent verification of the final submission package.

Checks (read-only, never modifies deliverables):
  * 10 YouTube Shorts: calendar mapping -> canonical clean file -> package copy (SHA-256 equal),
    ffprobe dimensions/duration/audio, no subtitle stream, no 'captioned' file in package.
  * Every calendar final_file_url / source_file_url exists inside the package.
  * Package contains no .env, __pycache__, .pyc, captioned shorts.
  * Official workbook hash unchanged; working copy keeps official formulas and data validations.

Run: py -3.12 -m engine.package_verify
"""
from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

OFFICIAL_SHA256 = "9142119c5633480b40e9de997883265cc930e4b868f2d128845fa9677b321b9a"
PKG_MAIN = Path("submission-package/main-output")
WORKBOOK_PKG = Path("submission-package/workbook/IncuBrix_Option2_Submission_Workbook_WORKING_COPY.xlsx")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def probe(path: Path) -> dict:
    out = subprocess.check_output(
        ["ffprobe", "-v", "error", "-show_entries",
         "stream=codec_type,codec_name,width,height,r_frame_rate:format=duration", "-of", "json", str(path)],
        text=True)
    data = json.loads(out)
    video = next((s for s in data["streams"] if s["codec_type"] == "video"), {})
    return {
        "width": video.get("width"), "height": video.get("height"),
        "fps": video.get("r_frame_rate"), "duration": round(float(data["format"]["duration"]), 2),
        "has_audio": any(s["codec_type"] == "audio" for s in data["streams"]),
        "has_subtitle_stream": any(s["codec_type"] == "subtitle" for s in data["streams"]),
    }


def split_paths(value: str) -> list[str]:
    return [p.strip() for p in value.split(";") if p.strip()]


def verify(root: str | Path = ".") -> dict:
    root = Path(root)
    cal_file = (root / "calendar_UPDATED.csv") if (root / "calendar_UPDATED.csv").exists() else (root / "calendar.csv")
    try:
        with cal_file.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    except UnicodeDecodeError:
        with cal_file.open(encoding="cp1252", newline="") as handle:
            rows = list(csv.DictReader(handle))
    problems: list[str] = []
    shorts = []
    for row in rows:
        if row["platform"] != "YouTube Shorts":
            continue
        rel = split_paths(row["final_file_url"])
        entry = {"asset_id": row["asset_id"], "mapped_file": rel[0] if rel else None}
        if len(rel) != 1:
            problems.append(f"{row['asset_id']}: expected exactly one final file, got {len(rel)}")
            shorts.append(entry)
            continue
        canon = root / rel[0]
        pkg = root / PKG_MAIN / rel[0]
        if "captioned" in rel[0].lower() or "youtube-shorts-captioned" in rel[0]:
            problems.append(f"{row['asset_id']}: mapped to a captioned file")
        if not canon.exists() or not pkg.exists():
            problems.append(f"{row['asset_id']}: missing canonical or package file")
            shorts.append(entry)
            continue
        info = probe(canon)
        entry.update(info)
        entry["sha256"] = sha256(canon)
        entry["package_sha256_equal"] = entry["sha256"] == sha256(pkg)
        entry["bytes"] = canon.stat().st_size
        if not entry["package_sha256_equal"]:
            problems.append(f"{row['asset_id']}: package copy differs from canonical")
        if (info["width"], info["height"]) != (720, 1280):
            problems.append(f"{row['asset_id']}: resolution {info['width']}x{info['height']}")
        if info["duration"] > 60 or not info["has_audio"] or info["has_subtitle_stream"]:
            problems.append(f"{row['asset_id']}: duration/audio/subtitle-stream check failed")
        shorts.append(entry)
    if len(shorts) != 10:
        problems.append(f"expected 10 shorts, found {len(shorts)}")

    # all mapped files exist in the package and match canonical SHA-256
    missing = []
    mismatched = []
    semicolon_fields = []
    for row in rows:
        for col in ("final_file_url", "source_file_url", "source_input"):
            if ";" in (row.get(col) or ""):
                semicolon_fields.append(f"{row['asset_id']}:{col}")
        for col in ("final_file_url", "source_file_url"):
            for rel in split_paths(row[col]):
                canon_f = root / rel
                pkg_f = root / PKG_MAIN / rel
                if not pkg_f.exists() or not canon_f.exists():
                    missing.append(f"{row['asset_id']}:{col}:{rel}")
                elif sha256(canon_f) != sha256(pkg_f):
                    mismatched.append(f"{row['asset_id']}:{col}:{rel}")
    if semicolon_fields:
        problems.append(f"semicolon-separated paths in calendar fields: {semicolon_fields[:5]}")
    for i in range(1, 11):
        aid = f"LI-{i:02d}"
        exp_dir = root / "content/exports/linkedin" / aid
        pkg_dir = root / PKG_MAIN / "content/exports/linkedin" / aid
        exp_pngs = sorted(exp_dir.glob("*.png")) if exp_dir.exists() else []
        pkg_pngs = sorted(pkg_dir.glob("*.png")) if pkg_dir.exists() else []
        if len(exp_pngs) < 6 or [p.name for p in exp_pngs] != [p.name for p in pkg_pngs]:
            problems.append(f"{aid}: export folder incomplete or package mismatch")
        else:
            for ep, pp in zip(exp_pngs, pkg_pngs):
                if sha256(ep) != sha256(pp):
                    mismatched.append(f"{aid}:export:{ep.name}")
    if missing:
        problems.append(f"{len(missing)} mapped files missing in package (first: {missing[:3]})")
    if mismatched:
        problems.append(f"{len(mismatched)} mapped files differ in SHA-256 in package (first: {mismatched[:3]})")
    if sha256(root / "calendar.csv") != sha256(root / PKG_MAIN / "calendar.csv"):
        problems.append("package calendar.csv differs from root calendar.csv")

    # forbidden content
    pkg_root = root / "submission-package"
    forbidden = [str(p) for p in pkg_root.rglob("*")
                 if p.name == ".env" or p.name == "__pycache__" or p.suffix == ".pyc"
                 or "captioned" in p.name.lower()]
    if forbidden:
        problems.append(f"forbidden files in package: {forbidden[:5]}")

    # workbook integrity
    official = root / "inputs" / "official_submission_workbook.xlsx"
    official_ok = sha256(official) == OFFICIAL_SHA256
    if not official_ok:
        problems.append("official workbook hash changed")
    formula_overwrites: list[str] = []
    validation_counts = {}
    wb_updated = root / WORKBOOK_PKG.with_name(WORKBOOK_PKG.stem + "_UPDATED" + WORKBOOK_PKG.suffix)
    if wb_updated.exists() and (not (root / WORKBOOK_PKG).exists() or wb_updated.stat().st_mtime >= (root / WORKBOOK_PKG).stat().st_mtime):
        wb_path = wb_updated
    else:
        wb_path = root / WORKBOOK_PKG
    if wb_path.exists():
        wo = openpyxl.load_workbook(official)
        wc = openpyxl.load_workbook(wb_path)
        for ws in wo.worksheets:
            other = wc[ws.title]
            validation_counts[ws.title] = (len(ws.data_validations.dataValidation),
                                           len(other.data_validations.dataValidation))
            for rw in ws.iter_rows():
                for cell in rw:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        if other[cell.coordinate].value != cell.value:
                            formula_overwrites.append(f"{ws.title}!{cell.coordinate}")
        wo.close()
        wc.close()
        if formula_overwrites:
            problems.append(f"formula cells changed: {formula_overwrites[:5]}")
        for title, (a, b) in validation_counts.items():
            if a != b:
                problems.append(f"data validation count differs on {title}: official {a} vs copy {b}")
    else:
        problems.append("working workbook missing")

    return {"status": "OK" if not problems else "PROBLEMS", "problems": problems, "shorts": shorts,
            "official_workbook_hash_ok": official_ok, "formula_overwrites": formula_overwrites,
            "data_validation_counts_official_vs_copy": validation_counts}


def main(argv: list[str] | None = None) -> int:
    result = verify(".")
    out = Path("outputs/final-verification.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "shorts"}, indent=2))
    for s in result["shorts"]:
        print(s.get("asset_id"), s.get("mapped_file"), f"{s.get('width')}x{s.get('height')}", s.get("duration"),
              s.get("sha256"), "pkg_equal=", s.get("package_sha256_equal"))
    return 0 if result["status"] == "OK" else 1


if __name__ == "__main__":
    sys.exit(main())
