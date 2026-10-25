"""Safely populate a copy of the official IncuBrix submission workbook."""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl


class WorkbookPopulationError(ValueError):
    """Raised when a calendar cannot safely populate the workbook."""


SHEET_NAME = "02_ASSETS"
HEADER_ROW = 4
FIRST_DATA_ROW = 5
LAST_DATA_ROW = 154
MAX_ASSETS = LAST_DATA_ROW - FIRST_DATA_ROW + 1
FORMULA_COLUMNS = {"duplicate_check", "completeness_check"}


def _read_calendar(csv_path: Path, expected_columns: list[str]) -> list[dict[str, str]]:
    try:
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != expected_columns:
                raise WorkbookPopulationError(
                    "Calendar CSV headers must exactly match the configured schema "
                    "in both names and order."
                )

            rows = []
            seen_asset_ids: set[str] = set()

            for row_number, row in enumerate(reader, start=2):
                if None in row or any(value is None for value in row.values()):
                    raise WorkbookPopulationError(
                        f"Calendar CSV row {row_number} has an invalid number of columns."
                    )

                asset_id = (row.get("asset_id") or "").strip()
                if not asset_id:
                    raise WorkbookPopulationError(
                        f"Calendar CSV row {row_number} has a blank asset_id."
                    )
                if asset_id in seen_asset_ids:
                    raise WorkbookPopulationError(
                        f"Calendar CSV row {row_number} duplicates asset_id "
                        f"{asset_id!r}."
                    )
                seen_asset_ids.add(asset_id)

                rows.append({key: value for key, value in row.items()})
    except OSError as exc:
        raise WorkbookPopulationError("Could not read the calendar CSV.") from exc
    except csv.Error as exc:
        raise WorkbookPopulationError("Calendar CSV is malformed.") from exc

    if len(rows) > MAX_ASSETS:
        raise WorkbookPopulationError(
            f"Calendar contains {len(rows)} assets; workbook limit is {MAX_ASSETS}."
        )

    return rows


def populate_workbook_copy(
    *,
    workbook_path: str | Path,
    csv_path: str | Path,
    schema_path: str | Path,
    output_path: str | Path,
) -> dict[str, object]:
    """Populate A:U in a new workbook copy, preserving template formulas and rules."""
    from .calendar_export import load_calendar_schema

    source = Path(workbook_path)
    calendar = Path(csv_path)
    schema_file = Path(schema_path)
    destination = Path(output_path)

    if not source.is_file():
        raise WorkbookPopulationError("Official workbook file does not exist.")
    if not calendar.is_file():
        raise WorkbookPopulationError("Calendar CSV file does not exist.")
    if source.resolve() == destination.resolve():
        raise WorkbookPopulationError("Output must not be the official source workbook.")
    if destination.exists():
        raise WorkbookPopulationError("Output already exists; refusing to overwrite it.")
    if destination.suffix.lower() != ".xlsx":
        raise WorkbookPopulationError("Output path must end in .xlsx.")

    schema = load_calendar_schema(schema_file)
    columns = schema["columns"]
    rows = _read_calendar(calendar, columns)

    try:
        workbook = openpyxl.load_workbook(source, data_only=False)
    except (OSError, ValueError) as exc:
        raise WorkbookPopulationError("Could not open the official workbook.") from exc

    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise WorkbookPopulationError(
                f"Official workbook is missing sheet {SHEET_NAME!r}."
            )

        sheet = workbook[SHEET_NAME]
        actual_headers = [
            sheet.cell(HEADER_ROW, col).value
            for col in range(1, len(columns) + 1)
        ]
        if actual_headers != columns:
            raise WorkbookPopulationError(
                "Official workbook asset headers do not match the configured schema."
            )

        allowed_values = schema.get("allowed_values", {})
        for row_number, row in enumerate(rows, start=1):
            for field, choices in allowed_values.items():
                value = row.get(field, "")
                if value and value not in choices:
                    raise WorkbookPopulationError(
                        f"CSV row {row_number}, field {field!r} has an unsupported value."
                    )

        input_columns = [
            (index, name)
            for index, name in enumerate(columns, start=1)
            if name not in FORMULA_COLUMNS
        ]
        if any(index > 21 for index, _ in input_columns):
            raise WorkbookPopulationError(
                "Unexpected input column after column U; refusing to write."
            )

        for offset, row in enumerate(rows):
            target_row = FIRST_DATA_ROW + offset
            for column_index, field in input_columns:
                sheet.cell(target_row, column_index).value = row[field]

        for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            if sheet.cell(row_number, 22).data_type != "f":
                raise WorkbookPopulationError(
                    f"Duplicate-check formula missing at V{row_number}."
                )
            if sheet.cell(row_number, 23).data_type != "f":
                raise WorkbookPopulationError(
                    f"Completeness formula missing at W{row_number}."
                )

        if len(sheet.data_validations.dataValidation) != 5:
            raise WorkbookPopulationError(
                "Expected five data validations; refusing to save."
            )
        if sheet.freeze_panes != "C5":
            raise WorkbookPopulationError(
                "Expected freeze panes C5; refusing to save."
            )

        if destination.exists():
            raise WorkbookPopulationError(
                "Output appeared during processing; refusing to overwrite it."
            )
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    except (OSError, ValueError) as exc:
        if isinstance(exc, WorkbookPopulationError):
            raise
        raise WorkbookPopulationError("Could not safely create the workbook copy.") from exc
    finally:
        workbook.close()

    try:
        check = openpyxl.load_workbook(destination, data_only=False, read_only=False)
        try:
            saved_sheet = check[SHEET_NAME]
            valid = (
                check.sheetnames == [
                    "00_START", "01_INSIGHT", "02_ASSETS",
                    "03_ENGINE_LAUNCH", "04_SELF_CHECK",
                ]
                and saved_sheet.freeze_panes == "C5"
                and len(saved_sheet.data_validations.dataValidation) == 5
                and all(
                    saved_sheet.cell(r, 22).data_type == "f"
                    and saved_sheet.cell(r, 23).data_type == "f"
                    for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)
                )
            )
        finally:
            check.close()

        if not valid:
            destination.unlink(missing_ok=True)
            raise WorkbookPopulationError(
                "Saved workbook failed structural verification; output was removed."
            )
    except (OSError, ValueError, KeyError) as exc:
        destination.unlink(missing_ok=True)
        if isinstance(exc, WorkbookPopulationError):
            raise
        raise WorkbookPopulationError(
            "Could not verify the saved workbook; output was removed."
        ) from exc

    return {
        "status": "CREATED",
        "output_path": str(destination.resolve()),
        "assets_written": len(rows),
        "asset_sheet": SHEET_NAME,
        "input_columns_written": 21,
        "formula_columns_preserved": ["duplicate_check", "completeness_check"],
        "original_workbook_modified": False,
    }
